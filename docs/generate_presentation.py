"""
Agentic SDLC System — PowerPoint Presentation Generator
Run: python docs/generate_presentation.py
Output: docs/Agentic_SDLC_System.pptx
Requires: pip install python-pptx
"""
import sys
import os

# Allow imports from project root
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

# ── Color palette ─────────────────────────────────────────────────────────────
DARK_BG    = RGBColor(0x1A, 0x1A, 0x2E)
ACCENT     = RGBColor(0x00, 0xD4, 0xFF)
ACCENT2    = RGBColor(0xFF, 0x6B, 0x35)
WHITE      = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_GRAY = RGBColor(0xCC, 0xCC, 0xCC)
GREEN      = RGBColor(0x00, 0xC8, 0x5A)
YELLOW     = RGBColor(0xFF, 0xD7, 0x00)
RED        = RGBColor(0xFF, 0x4C, 0x4C)
CARD_BG    = RGBColor(0x16, 0x21, 0x3E)

prs = Presentation()
prs.slide_width  = Inches(13.33)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


# ── Shared helpers ────────────────────────────────────────────────────────────
def slide():
    return prs.slides.add_slide(BLANK)


def bg(sl, color=DARK_BG):
    sh = sl.shapes.add_shape(1, 0, 0, prs.slide_width, prs.slide_height)
    sh.fill.solid()
    sh.fill.fore_color.rgb = color
    sh.line.fill.background()


def box(sl, x, y, w, h, fill=CARD_BG, line_color=None, line_width=Pt(0)):
    sh = sl.shapes.add_shape(1, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    if line_color:
        sh.line.color.rgb = line_color
        sh.line.width = line_width
    else:
        sh.line.fill.background()
    return sh


def txt(sl, text, x, y, w, h, size=18, bold=False, color=WHITE,
        align=PP_ALIGN.LEFT, wrap=True):
    tb = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tb.word_wrap = wrap
    tf = tb.text_frame
    tf.word_wrap = wrap
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return tb


def accent_bar(sl, y=0.55, color=ACCENT):
    sh = sl.shapes.add_shape(1, 0, Inches(y), prs.slide_width, Inches(0.06))
    sh.fill.solid()
    sh.fill.fore_color.rgb = color
    sh.line.fill.background()


helpers = (slide, bg, box, txt, accent_bar)

# ── Import and run all slide builders ────────────────────────────────────────
from docs.generate_presentation_p1 import *   # noqa: F401,F403 — re-runs slide creation
from docs.generate_presentation_p2 import add_slides_6_to_10
from docs.generate_presentation_p3 import add_slides_11_to_14

# p1 already added slides 1-5 to `prs` via module-level code — rebuild cleanly
# Reset and rebuild all slides in order
prs2 = Presentation()
prs2.slide_width  = Inches(13.33)
prs2.slide_height = Inches(7.5)
BLANK2 = prs2.slide_layouts[6]


def slide2():
    return prs2.slides.add_slide(BLANK2)


def bg2(sl, color=DARK_BG):
    sh = sl.shapes.add_shape(1, 0, 0, prs2.slide_width, prs2.slide_height)
    sh.fill.solid()
    sh.fill.fore_color.rgb = color
    sh.line.fill.background()


def box2(sl, x, y, w, h, fill=CARD_BG, line_color=None, line_width=Pt(0)):
    sh = sl.shapes.add_shape(1, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    if line_color:
        sh.line.color.rgb = line_color
        sh.line.width = line_width
    else:
        sh.line.fill.background()
    return sh


def txt2(sl, text, x, y, w, h, size=18, bold=False, color=WHITE,
         align=PP_ALIGN.LEFT, wrap=True):
    tb = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tb.word_wrap = wrap
    tf = tb.text_frame
    tf.word_wrap = wrap
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return tb


def accent_bar2(sl, y=0.55, color=ACCENT):
    sh = sl.shapes.add_shape(1, 0, Inches(y), prs2.slide_width, Inches(0.06))
    sh.fill.solid()
    sh.fill.fore_color.rgb = color
    sh.line.fill.background()


h2 = (slide2, bg2, box2, txt2, accent_bar2)


def build_slides_1_to_5():
    """Rebuild slides 1-5 using prs2."""
    # Slide 1 — Title
    sl = slide2()
    bg2(sl)
    accent_bar2(sl, y=0.0, color=ACCENT)
    accent_bar2(sl, y=7.44, color=ACCENT)
    box2(sl, 0, 0, 0.08, 7.5, fill=ACCENT)
    txt2(sl, "AGENTIC SDLC SYSTEM", 0.3, 1.6, 12.5, 1.0,
         size=44, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    txt2(sl, "Multi-Agent Workflow Orchestration Across the Full Software Development Lifecycle",
         0.5, 2.7, 12.3, 0.8, size=20, color=ACCENT, align=PP_ALIGN.CENTER)
    box2(sl, 3.5, 3.7, 6.3, 0.06, fill=ACCENT)
    txt2(sl, "From Requirement to Reviewable Engineering Outcome — Automatically",
         0.5, 3.9, 12.3, 0.6, size=16, color=LIGHT_GRAY, align=PP_ALIGN.CENTER)
    txt2(sl, "Greenfield  |  Brownfield  |  Ambiguous Requirements",
         0.5, 4.7, 12.3, 0.5, size=15, color=ACCENT2, align=PP_ALIGN.CENTER)
    txt2(sl, "2025", 0.5, 6.8, 12.3, 0.4, size=13,
         color=LIGHT_GRAY, align=PP_ALIGN.CENTER)

    # Slide 2 — Problem
    sl = slide2()
    bg2(sl)
    accent_bar2(sl, y=0.55)
    txt2(sl, "THE PROBLEM", 0.4, 0.1, 12.5, 0.5, size=28, bold=True, color=WHITE)
    problems = [
        ("Requirements are ambiguous",
         "Engineers waste hours clarifying vague specs before writing a single line of code."),
        ("SDLC steps are disconnected",
         "Architecture, code, tests, and docs are produced in silos with no shared context."),
        ("No controlled autonomy",
         "Either fully manual (slow) or fully automated (risky). No middle ground."),
        ("Brownfield work is opaque",
         "Changing existing systems without understanding impact causes regressions."),
    ]
    for i, (title, desc) in enumerate(problems):
        col = i % 2
        row = i // 2
        bx = 0.3 + col * 6.5
        by = 1.0 + row * 2.8
        box2(sl, bx, by, 6.2, 2.5, fill=CARD_BG, line_color=ACCENT2, line_width=Pt(1.5))
        txt2(sl, f"  {title}", bx + 0.15, by + 0.15, 5.9, 0.45,
             size=16, bold=True, color=ACCENT2)
        txt2(sl, desc, bx + 0.2, by + 0.65, 5.8, 1.6, size=13, color=LIGHT_GRAY)

    # Slide 3 — Solution
    sl = slide2()
    bg2(sl)
    accent_bar2(sl, y=0.55)
    txt2(sl, "THE SOLUTION", 0.4, 0.1, 12.5, 0.5, size=28, bold=True, color=WHITE)
    txt2(sl, "One requirement in. Complete engineering outcome out.",
         0.4, 0.65, 12.5, 0.45, size=17, color=ACCENT, bold=True)
    steps = [
        ("1", "Understand", "Classify scenario\nResolve ambiguities\nExtract constraints"),
        ("2", "Decompose", "Break into tasks\nDefine dependencies\nOrder execution"),
        ("3", "Execute", "9 specialised agents\nApproval gates\nError recovery"),
        ("4", "Generate", "Code + API + Schema\nTests + Docs\nDocker artifacts"),
        ("5", "Validate", "Risk identification\nStatic checks\nTest strategy"),
    ]
    for i, (num, title, desc) in enumerate(steps):
        bx = 0.3 + i * 2.55
        box2(sl, bx, 1.3, 2.3, 2.9, fill=CARD_BG, line_color=ACCENT, line_width=Pt(1.5))
        box2(sl, bx + 0.75, 1.35, 0.8, 0.7, fill=ACCENT)
        txt2(sl, num, bx + 0.75, 1.35, 0.8, 0.7,
             size=22, bold=True, color=DARK_BG, align=PP_ALIGN.CENTER)
        txt2(sl, title, bx + 0.1, 2.15, 2.1, 0.4,
             size=15, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        txt2(sl, desc, bx + 0.15, 2.6, 2.0, 1.4,
             size=11, color=LIGHT_GRAY, align=PP_ALIGN.CENTER)
        if i < 4:
            txt2(sl, "->", bx + 2.3, 2.5, 0.25, 0.4,
                 size=18, bold=True, color=ACCENT, align=PP_ALIGN.CENTER)
    txt2(sl, "Human-in-the-Loop approval gates at Architecture, Code Generation, and Validation",
         0.4, 4.5, 12.5, 0.5, size=14, color=ACCENT2, align=PP_ALIGN.CENTER)

    # Slide 4 — Architecture
    sl = slide2()
    bg2(sl)
    accent_bar2(sl, y=0.55)
    txt2(sl, "SYSTEM ARCHITECTURE", 0.4, 0.1, 12.5, 0.5, size=28, bold=True, color=WHITE)
    components = [
        ("CLI / main.py", "Entry point, structured JSON logging,\n/health + /metrics HTTP endpoints"),
        ("WorkflowOrchestrator", "Dependency scheduler, approval gates,\nretry logic, failure propagation"),
        ("WorkflowState", "Shared mutable blackboard — single\nsource of truth across all agents"),
        ("OutputWriter", "Path-safe artifact persistence\nwith traversal protection"),
    ]
    for i, (name, desc) in enumerate(components):
        by = 0.85 + i * 1.55
        box2(sl, 0.3, by, 4.5, 1.4, fill=CARD_BG, line_color=ACCENT, line_width=Pt(1))
        txt2(sl, name, 0.45, by + 0.1, 4.2, 0.4, size=13, bold=True, color=ACCENT)
        txt2(sl, desc, 0.45, by + 0.55, 4.2, 0.75, size=11, color=LIGHT_GRAY)
    agents = ["RequirementAgent", "DecompositionAgent", "ArchitectureAgent",
              "SchemaAgent", "CodeGenerationAgent", "TestGenerationAgent",
              "ValidationAgent", "DocumentationAgent", "CodebaseReasoningAgent"]
    box2(sl, 5.2, 0.85, 3.8, 6.3, fill=CARD_BG)
    txt2(sl, "AGENT REGISTRY", 5.35, 0.9, 3.5, 0.4, size=13, bold=True, color=ACCENT)
    for i, a in enumerate(agents):
        txt2(sl, f"  {a}", 5.35, 1.4 + i * 0.6, 3.5, 0.4, size=11, color=WHITE)
    decisions = [
        "Shared WorkflowState — no message passing",
        "Dependency-aware scheduler",
        "Retry (max 2) + skip-dependents on failure",
        "html.escape on all user input",
        "Path confinement in OutputWriter",
        "auto_approve flag for CI/CD",
    ]
    box2(sl, 9.2, 0.85, 3.9, 6.3, fill=CARD_BG)
    txt2(sl, "KEY DECISIONS", 9.35, 0.9, 3.6, 0.4, size=13, bold=True, color=ACCENT2)
    for i, d in enumerate(decisions):
        txt2(sl, f"  {d}", 9.35, 1.4 + i * 0.85, 3.6, 0.75, size=10, color=LIGHT_GRAY)

    # Slide 5 — Agent Deep Dive
    sl = slide2()
    bg2(sl)
    accent_bar2(sl, y=0.55)
    txt2(sl, "AGENT DEEP DIVE", 0.4, 0.1, 12.5, 0.5, size=28, bold=True, color=WHITE)
    agents_detail = [
        ("RequirementAgent", ACCENT,
         "Classify: greenfield/brownfield/ambiguous\n"
         "Detect ambiguities -> auto-resolve\n"
         "Normalize + extract success criteria\n"
         "html.escape all user input"),
        ("DecompositionAgent", ACCENT2,
         "Greenfield: 6-7 tasks with analytics\n"
         "Brownfield: codebase_analysis first\n"
         "Explicit depends_on for every task\n"
         "Injects tasks into WorkflowState"),
        ("ArchitectureAgent", GREEN,
         "System components + data flows\n"
         "Cache-aside pattern with Redis\n"
         "Async analytics pipeline design\n"
         "AWS deployment recommendations"),
        ("CodeGenerationAgent", YELLOW,
         "10 FastAPI source files generated\n"
         "Layered: routes->service->repo->DB\n"
         "Cache-first redirect (<1ms hot)\n"
         "Fire-and-forget analytics"),
        ("ValidationAgent", RED,
         "6 risks with mitigations\n"
         "Static checks on all artifacts\n"
         "Blocking vs non-blocking checks\n"
         "Full test strategy defined"),
        ("CodebaseReasoningAgent", ACCENT,
         "Brownfield only: walks directory\n"
         "Scores file relevance by keywords\n"
         "Detects API/schema breaking risk\n"
         "Path-confined to base directory"),
    ]
    for i, (name, color, detail) in enumerate(agents_detail):
        col = i % 3
        row = i // 3
        bx = 0.3 + col * 4.35
        by = 0.85 + row * 3.1
        box2(sl, bx, by, 4.1, 2.85, fill=CARD_BG, line_color=color, line_width=Pt(1.5))
        txt2(sl, name, bx + 0.15, by + 0.1, 3.8, 0.4, size=13, bold=True, color=color)
        txt2(sl, detail, bx + 0.15, by + 0.55, 3.8, 2.1, size=10, color=LIGHT_GRAY)


build_slides_1_to_5()
add_slides_6_to_10(prs2, slide2, bg2, box2, txt2, accent_bar2)
add_slides_11_to_14(prs2, slide2, bg2, box2, txt2, accent_bar2)

out_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Agentic_SDLC_System.pptx")
prs2.save(out_path)
print(f"Presentation saved: {out_path}")
print(f"Total slides: {len(prs2.slides)}")
