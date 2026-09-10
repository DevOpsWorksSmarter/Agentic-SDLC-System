"""
Generate Agentic SDLC System PowerPoint presentation.
Run: python docs/generate_presentation.py
Requires: pip install python-pptx
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
import os

# ── Color palette ─────────────────────────────────────────────────────────────
DARK_BG    = RGBColor(0x1A, 0x1A, 0x2E)   # deep navy
ACCENT     = RGBColor(0x00, 0xD4, 0xFF)   # cyan
ACCENT2    = RGBColor(0xFF, 0x6B, 0x35)   # orange
WHITE      = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_GRAY = RGBColor(0xCC, 0xCC, 0xCC)
GREEN      = RGBColor(0x00, 0xC8, 0x5A)
YELLOW     = RGBColor(0xFF, 0xD7, 0x00)
RED        = RGBColor(0xFF, 0x4C, 0x4C)
CARD_BG    = RGBColor(0x16, 0x21, 0x3E)

prs = Presentation()
prs.slide_width  = Inches(13.33)
prs.slide_height = Inches(7.5)

BLANK = prs.slide_layouts[6]  # completely blank


def slide():
    return prs.slides.add_slide(BLANK)


def bg(sl, color=DARK_BG):
    sh = sl.shapes.add_shape(1, 0, 0, prs.slide_width, prs.slide_height)
    sh.fill.solid()
    sh.fill.fore_color.rgb = color
    sh.line.fill.background()


def box(sl, x, y, w, h, fill=CARD_BG, line_color=None, line_width=Pt(0)):
    sh = sl.shapes.add_shape(1, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    if line_color:
        sh.line.color.rgb = line_color
        sh.line.width = line_width
    else:
        sh.line.fill.background()
    return sh


def txt(sl, text, x, y, w, h, size=18, bold=False, color=WHITE,
        align=PP_ALIGN.LEFT, wrap=True):
    tb = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tb.word_wrap = wrap
    tf = tb.text_frame
    tf.word_wrap = wrap
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return tb


def accent_bar(sl, y=0.55, color=ACCENT):
    sh = sl.shapes.add_shape(1, Inches(0), Inches(y), prs.slide_width, Inches(0.06))
    sh.fill.solid()
    sh.fill.fore_color.rgb = color
    sh.line.fill.background()


def bullet_box(sl, items, x, y, w, h, title=None, title_color=ACCENT,
               item_size=14, title_size=16):
    box(sl, x, y, w, h)
    ty = y + 0.12
    if title:
        txt(sl, title, x + 0.15, ty, w - 0.3, 0.35,
            size=title_size, bold=True, color=title_color)
        ty += 0.38
    for item in items:
        txt(sl, item, x + 0.2, ty, w - 0.35, 0.32, size=item_size, color=WHITE)
        ty += 0.3


# ─────────────────────────────────────────────────────────────────────────────
# SLIDE 1 — Title
# ─────────────────────────────────────────────────────────────────────────────
sl = slide()
bg(sl)
accent_bar(sl, y=0.0, color=ACCENT)
accent_bar(sl, y=7.44, color=ACCENT)

# Gradient-like side bar
box(sl, 0, 0, 0.08, 7.5, fill=ACCENT)

txt(sl, "AGENTIC SDLC SYSTEM", 0.3, 1.6, 12.5, 1.0,
    size=44, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
txt(sl, "Multi-Agent Workflow Orchestration Across the Full Software Development Lifecycle",
    0.5, 2.7, 12.3, 0.8, size=20, color=ACCENT, align=PP_ALIGN.CENTER)

box(sl, 3.5, 3.7, 6.3, 0.06, fill=ACCENT)

txt(sl, "From Requirement to Reviewable Engineering Outcome — Automatically",
    0.5, 3.9, 12.3, 0.6, size=16, color=LIGHT_GRAY, align=PP_ALIGN.CENTER)

txt(sl, "Greenfield  |  Brownfield  |  Ambiguous Requirements",
    0.5, 4.7, 12.3, 0.5, size=15, color=ACCENT2, align=PP_ALIGN.CENTER)

txt(sl, "2025", 0.5, 6.8, 12.3, 0.4, size=13,
    color=LIGHT_GRAY, align=PP_ALIGN.CENTER)


# ─────────────────────────────────────────────────────────────────────────────
# SLIDE 2 — Problem Statement
# ─────────────────────────────────────────────────────────────────────────────
sl = slide()
bg(sl)
accent_bar(sl, y=0.55)
txt(sl, "THE PROBLEM", 0.4, 0.1, 12.5, 0.5,
    size=28, bold=True, color=WHITE, align=PP_ALIGN.LEFT)

problems = [
    ("Requirements are ambiguous", "Engineers waste hours clarifying vague specs before writing a single line of code."),
    ("SDLC steps are disconnected", "Architecture, code, tests, and docs are produced in silos with no shared context."),
    ("No controlled autonomy", "Either fully manual (slow) or fully automated (risky). No middle ground."),
    ("Brownfield work is opaque", "Changing existing systems without understanding impact causes regressions."),
]

for i, (title, desc) in enumerate(problems):
    col = i % 2
    row = i // 2
    bx = 0.3 + col * 6.5
    by = 1.0 + row * 2.8
    box(sl, bx, by, 6.2, 2.5, fill=CARD_BG, line_color=ACCENT2, line_width=Pt(1.5))
    txt(sl, f"  {title}", bx + 0.15, by + 0.15, 5.9, 0.45,
        size=16, bold=True, color=ACCENT2)
    txt(sl, desc, bx + 0.2, by + 0.65, 5.8, 1.6, size=13, color=LIGHT_GRAY)


# ─────────────────────────────────────────────────────────────────────────────
# SLIDE 3 — Solution Overview
# ─────────────────────────────────────────────────────────────────────────────
sl = slide()
bg(sl)
accent_bar(sl, y=0.55)
txt(sl, "THE SOLUTION", 0.4, 0.1, 12.5, 0.5,
    size=28, bold=True, color=WHITE)

txt(sl, "One requirement in. Complete engineering outcome out.",
    0.4, 0.65, 12.5, 0.45, size=17, color=ACCENT, bold=True)

steps = [
    ("1", "Understand", "Classify scenario\nResolve ambiguities\nExtract constraints"),
    ("2", "Decompose", "Break into tasks\nDefine dependencies\nOrder execution"),
    ("3", "Execute", "9 specialised agents\nApproval gates\nError recovery"),
    ("4", "Generate", "Code + API + Schema\nTests + Docs\nDocker artifacts"),
    ("5", "Validate", "Risk identification\nStatic checks\nTest strategy"),
]

arrow_y = 3.5
for i, (num, title, desc) in enumerate(steps):
    bx = 0.3 + i * 2.55
    box(sl, bx, 1.3, 2.3, 2.9, fill=CARD_BG, line_color=ACCENT, line_width=Pt(1.5))
    box(sl, bx + 0.75, 1.35, 0.8, 0.7, fill=ACCENT)
    txt(sl, num, bx + 0.75, 1.35, 0.8, 0.7,
        size=22, bold=True, color=DARK_BG, align=PP_ALIGN.CENTER)
    txt(sl, title, bx + 0.1, 2.15, 2.1, 0.4,
        size=15, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    txt(sl, desc, bx + 0.15, 2.6, 2.0, 1.4,
        size=11, color=LIGHT_GRAY, align=PP_ALIGN.CENTER)
    if i < 4:
        txt(sl, "->", bx + 2.3, 2.5, 0.25, 0.4,
            size=18, bold=True, color=ACCENT, align=PP_ALIGN.CENTER)

txt(sl, "Human-in-the-Loop approval gates at Architecture, Code Generation, and Validation",
    0.4, 4.5, 12.5, 0.5, size=14, color=ACCENT2, align=PP_ALIGN.CENTER)
box(sl, 1.5, 4.45, 10.3, 0.06, fill=ACCENT2)


# ─────────────────────────────────────────────────────────────────────────────
# SLIDE 4 — System Architecture
# ─────────────────────────────────────────────────────────────────────────────
sl = slide()
bg(sl)
accent_bar(sl, y=0.55)
txt(sl, "SYSTEM ARCHITECTURE", 0.4, 0.1, 12.5, 0.5,
    size=28, bold=True, color=WHITE)

# Left column: components
components = [
    ("CLI / main.py", "Entry point, structured JSON logging,\n/health + /metrics HTTP endpoints"),
    ("WorkflowOrchestrator", "Dependency scheduler, approval gates,\nretry logic, failure propagation"),
    ("WorkflowState", "Shared mutable blackboard — single\nsource of truth across all agents"),
    ("OutputWriter", "Path-safe artifact persistence\nwith traversal protection"),
]
for i, (name, desc) in enumerate(components):
    by = 0.85 + i * 1.55
    box(sl, 0.3, by, 4.5, 1.4, fill=CARD_BG, line_color=ACCENT, line_width=Pt(1))
    txt(sl, name, 0.45, by + 0.1, 4.2, 0.4, size=13, bold=True, color=ACCENT)
    txt(sl, desc, 0.45, by + 0.55, 4.2, 0.75, size=11, color=LIGHT_GRAY)

# Right column: agent registry
agents = [
    "RequirementAgent",
    "DecompositionAgent",
    "ArchitectureAgent",
    "SchemaAgent",
    "CodeGenerationAgent",
    "TestGenerationAgent",
    "ValidationAgent",
    "DocumentationAgent",
    "CodebaseReasoningAgent",
]
box(sl, 5.2, 0.85, 3.8, 6.3, fill=CARD_BG)
txt(sl, "AGENT REGISTRY", 5.35, 0.9, 3.5, 0.4,
    size=13, bold=True, color=ACCENT)
for i, a in enumerate(agents):
    txt(sl, f"  {a}", 5.35, 1.4 + i * 0.6, 3.5, 0.4, size=11, color=WHITE)

# Key design decisions
decisions = [
    "Shared WorkflowState — no message passing overhead",
    "Dependency-aware scheduler — tasks run when deps are terminal",
    "Retry (max 2) + skip-dependents on final failure",
    "html.escape on all user input — XSS prevention",
    "Path confinement in OutputWriter — traversal blocked",
    "auto_approve flag — enables full CI/CD execution",
]
box(sl, 9.2, 0.85, 3.9, 6.3, fill=CARD_BG)
txt(sl, "KEY DECISIONS", 9.35, 0.9, 3.6, 0.4,
    size=13, bold=True, color=ACCENT2)
for i, d in enumerate(decisions):
    txt(sl, f"  {d}", 9.35, 1.4 + i * 0.85, 3.6, 0.75, size=10, color=LIGHT_GRAY)


# ─────────────────────────────────────────────────────────────────────────────
# SLIDE 5 — Agent Deep Dive
# ─────────────────────────────────────────────────────────────────────────────
sl = slide()
bg(sl)
accent_bar(sl, y=0.55)
txt(sl, "AGENT DEEP DIVE", 0.4, 0.1, 12.5, 0.5,
    size=28, bold=True, color=WHITE)

agents_detail = [
    ("RequirementAgent", ACCENT,
     "Classify: greenfield / brownfield / ambiguous\n"
     "Detect ambiguities -> auto-resolve with defaults\n"
     "Normalize + extract success criteria\n"
     "html.escape all user input"),
    ("DecompositionAgent", ACCENT2,
     "Greenfield: 6-7 tasks with analytics injection\n"
     "Brownfield: codebase_analysis -> impact_assessment\n"
     "Explicit depends_on for every task\n"
     "Injects tasks into shared WorkflowState"),
    ("ArchitectureAgent", GREEN,
     "System components + data flows + infra topology\n"
     "Cache-aside pattern with Redis\n"
     "Async analytics pipeline design\n"
     "AWS deployment recommendations"),
    ("CodeGenerationAgent", YELLOW,
     "10 FastAPI source files generated\n"
     "Layered: routes -> service -> repository -> DB\n"
     "Cache-first redirect path (<1ms hot)\n"
     "Fire-and-forget analytics (never blocks redirect)"),
    ("ValidationAgent", RED,
     "6 risks identified with mitigations\n"
     "Static checks on all artifact presence\n"
     "Blocking vs non-blocking check classification\n"
     "Full test strategy (unit/integration/perf/chaos)"),
    ("CodebaseReasoningAgent", ACCENT,
     "Brownfield only: walks directory tree\n"
     "Scores file relevance by keyword matching\n"
     "Detects API/schema breaking change risk\n"
     "Path-confined to declared base directory"),
]

for i, (name, color, detail) in enumerate(agents_detail):
    col = i % 3
    row = i // 3
    bx = 0.3 + col * 4.35
    by = 0.85 + row * 3.1
    box(sl, bx, by, 4.1, 2.85, fill=CARD_BG, line_color=color, line_width=Pt(1.5))
    txt(sl, name, bx + 0.15, by + 0.1, 3.8, 0.4,
        size=13, bold=True, color=color)
    txt(sl, detail, bx + 0.15, by + 0.55, 3.8, 2.1, size=10, color=LIGHT_GRAY)

print("Slides 1-5 done")
