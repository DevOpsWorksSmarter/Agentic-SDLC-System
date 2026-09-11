from pptx import Presentation
from pptx.util import Inches, Pt

prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)

def slide(title, bullets):
    s=prs.slides.add_slide(prs.slide_layouts[5])
    tb=s.shapes.add_textbox(Inches(.7), Inches(.45), Inches(12), Inches(.7))
    p=tb.text_frame.paragraphs[0]; p.text=title; p.font.size=Pt(28); p.font.bold=True
    body=s.shapes.add_textbox(Inches(.9), Inches(1.4), Inches(11.6), Inches(5.4))
    tf=body.text_frame; tf.word_wrap=True
    for i,b in enumerate(bullets):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph(); p.text=b; p.font.size=Pt(20); p.space_after=Pt(12)
    return s

slide('Agentic SDLC System — V7', ['Production-oriented agentic workflow for transforming requirements into reviewable engineering outcomes','Mandatory use case: scalable URL shortener with APIs, persistence and analytics','Goal: strong evidence across AI reasoning, SRE, security, validation and controlled autonomy'])
slide('Assignment alignment', ['Requirement understanding and ambiguity detection','Structured task decomposition and dependency-aware orchestration','Greenfield + brownfield + ambiguous scenarios','Production-quality code, API contracts, tests and documentation','Validation, risk management, guardrails and human approval'])
slide('End-to-end execution model', ['Requirement → AI reasoning → task graph → architecture/schema → code/tests','Validation → security review + SRE review + engineering review','Deterministic release gate → human approval → final evidence package','Shared WorkflowState provides cross-step coordination'])
slide('AI / agentic design', ['Optional OpenAI-compatible structured JSON adapter','Schema-constrained outputs and deterministic fallbacks','AI is advisory: policy and validation remain deterministic','Bounded repair agent prevents arbitrary autonomous mutation','Evidence records preserve reasoning and decisions'])
slide('SRE design', ['SLO: 99.9% monthly availability','Redirect p99 target <50ms warm-cache; create p99 <200ms','10k RPS target with stateless replicas, Redis and PostgreSQL HA','Resilience: Redis/DB failure, analytics isolation, duplicate concurrency','Metrics, structured logs, traces and operational runbook'])
slide('Security and governance', ['Input validation and safe output handling','Rate limiting, HTTPS, non-root containers, read-only filesystem','Dependency/security checks and threat review','High-impact tasks require approval','Release is blocked when required evidence is missing'])
slide('Scenario coverage', ['Greenfield: full URL-shortener lifecycle','Brownfield: codebase analysis, impact assessment and regression path','Ambiguous: explicit assumptions for missing performance/scale targets','Failure path: retry → diagnosis/repair → revalidation','All paths converge on validation and release evidence'])
slide('Evidence package', ['Requirement and assumptions','Task/dependency plan','Architecture and API/schema artifacts','Generated code and tests','Validation, security, SRE and engineering reviews','Release decision, audit log and engineering summary'])
slide('Verification', ['26 automated system tests passing in the upgraded repository','Mandatory URL-shortener workflow completed end-to-end','Deterministic offline mode keeps CI reproducible','LLM mode can be enabled through environment configuration'])
slide('Production evolution', ['Replace local workflow state with durable orchestration','Use managed PostgreSQL/Redis and distributed rate limiting','Integrate enterprise IAM, approvals and secret manager','Add OpenTelemetry, SBOM/signing and policy admission controls','Connect deployment contracts to EKS/ECS or target platform'])
slide('Interview positioning', ['This is not a generic chatbot: it is a governed SDLC workflow','AI proposes engineering decisions; deterministic controls validate them','System demonstrates ownership of quality, failure handling and operational readiness','The strongest differentiator is the complete evidence chain from requirement to release'])

out='docs/Agentic_SDLC_System_V7_Presentation.pptx'; prs.save(out); print(out)
